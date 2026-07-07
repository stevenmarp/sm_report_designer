# -*- coding: utf-8 -*-

import base64
import logging
import os
import platform
import re
import shutil
import struct
import subprocess
import tempfile
import zipfile
from copy import deepcopy
from io import BytesIO

from lxml import etree

from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError
from odoo.tools.misc import formatLang, format_date, format_datetime
from odoo.tools.safe_eval import safe_eval

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None

_logger = logging.getLogger(__name__)

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_NS = 'http://www.w3.org/XML/1998/namespace'
WP_NS = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
PIC_NS = 'http://schemas.openxmlformats.org/drawingml/2006/picture'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'
CT_NS = 'http://schemas.openxmlformats.org/package/2006/content-types'
IMG_REL_TYPE = ('http://schemas.openxmlformats.org/officeDocument/'
                '2006/relationships/image')


# ---------------------------------------------------------------------------
#  Field mapping (child model)
# ---------------------------------------------------------------------------

class SmReportTemplateField(models.Model):
    _name = 'sm.report.template.field'
    _description = 'Report Template Field Mapping'
    _order = 'sequence, id'

    template_id = fields.Many2one(
        'sm.report.template', string='Template',
        required=True, ondelete='cascade',
    )
    placeholder = fields.Char('Placeholder', required=True)
    field_path = fields.Char(
        'Field Path',
        help='Dot-separated path to a field on the model, e.g. partner_id.name',
    )
    is_loop = fields.Boolean(
        'Is Loop',
        help='Check if this placeholder marks a repeating section (One2many field).',
    )
    loop_parent = fields.Char(
        'Loop Parent',
        help='If this placeholder is inside a loop, enter the loop placeholder name.',
    )
    auto_resolved = fields.Boolean('Auto Resolved', readonly=True)
    description = fields.Char('Description')
    sequence = fields.Integer('Sequence', default=10)


# ---------------------------------------------------------------------------
#  Report template (parent model)
# ---------------------------------------------------------------------------

class SmReportTemplate(models.Model):
    _name = 'sm.report.template'
    _description = 'Report Template Designer'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'name'

    name = fields.Char('Template Name', required=True, tracking=True)
    model_id = fields.Many2one(
        'ir.model', string='Odoo Model', required=True,
        ondelete='cascade', tracking=True,
        domain=[('transient', '=', False)],
        help='Select the Odoo model this template applies to.',
    )
    model_name = fields.Char(related='model_id.model', store=True, readonly=True)

    template_file = fields.Binary(
        'Template File', required=True, attachment=True,
    )
    template_filename = fields.Char('File Name')
    template_type = fields.Selection(
        [('docx', 'Word (.docx)'), ('xlsx', 'Excel (.xlsx)')],
        string='Type', compute='_compute_template_type', store=True,
    )
    output_format = fields.Selection(
        [
            ('original', 'Same as Template'),
            ('pdf', 'PDF'),
            ('html', 'HTML'),
            ('odt', 'ODT (OpenDocument)'),
            ('rtf', 'RTF (Rich Text)'),
        ],
        string='Output Format', default='pdf', required=True,
    )
    merge_pdf = fields.Boolean(
        'Merge into Single PDF',
        default=True,
        help='When printing multiple records as PDF, merge all pages '
             'into a single PDF file instead of a ZIP archive.',
    )

    # Cover page
    cover_file = fields.Binary(
        'Cover Page Template', attachment=True,
        help='Optional DOCX cover page. Rendered with the same context '
             'and prepended to the output PDF.',
    )
    cover_filename = fields.Char('Cover File Name')

    # Domain-based template selection
    use_domain = fields.Boolean(
        'Domain Filter',
        help='When enabled, this template is only available for records '
             'that match the specified domain.',
    )
    template_domain = fields.Char(
        'Domain',
        default='[]',
        help='Odoo domain expression to filter which records can use this template.\n'
             'Example: [("state", "=", "sale")] to only match confirmed sales.',
    )
    priority = fields.Integer(
        'Priority', default=10,
        help='When multiple templates match a record, the one with the '
             'lowest priority number is used first.',
    )

    field_ids = fields.One2many(
        'sm.report.template.field', 'template_id', string='Field Mappings',
    )
    state = fields.Selection(
        [('draft', 'Draft'), ('ready', 'Ready'), ('active', 'Active')],
        default='draft', tracking=True, copy=False,
    )
    active = fields.Boolean(default=True)
    company_id = fields.Many2one(
        'res.company', string='Company',
        default=lambda self: self.env.company,
    )
    server_action_id = fields.Many2one(
        'ir.actions.server', string='Print Action',
        readonly=True, copy=False,
    )

    # ------------------------------------------------------------------
    #  Computed
    # ------------------------------------------------------------------

    @api.depends('template_filename')
    def _compute_template_type(self):
        for rec in self:
            if rec.template_filename:
                ext = rec.template_filename.rsplit('.', 1)[-1].lower()
                rec.template_type = ext if ext in ('docx', 'xlsx') else False
            else:
                rec.template_type = False

    # ------------------------------------------------------------------
    #  Actions (buttons)
    # ------------------------------------------------------------------

    def action_parse_template(self):
        """Parse the uploaded template and extract placeholders."""
        self.ensure_one()
        if not self.template_file:
            raise UserError(_('Please upload a template file first.'))
        if not self.model_id:
            raise UserError(_('Please select an Odoo model first.'))

        file_bytes = base64.b64decode(self.template_file)

        if self.template_type == 'docx':
            simple, loops, loop_children = self._parse_docx_placeholders(file_bytes)
        elif self.template_type == 'xlsx':
            simple, loops, loop_children = self._parse_xlsx_placeholders(file_bytes)
        else:
            raise UserError(_(
                'Unsupported file type. Please upload a .docx or .xlsx file.'
            ))

        # Remove old mappings
        self.field_ids.unlink()

        seq = 0

        # Create loop mappings
        for loop_name in sorted(loops):
            seq += 10
            field_path, resolved = self._auto_resolve_placeholder(loop_name)
            self.env['sm.report.template.field'].create({
                'template_id': self.id,
                'placeholder': loop_name,
                'field_path': field_path,
                'is_loop': True,
                'auto_resolved': resolved,
                'sequence': seq,
                'description': _('Loop over %s') % loop_name,
            })

        # Create regular placeholder mappings
        for ph in sorted(simple):
            seq += 10
            field_path, resolved = self._auto_resolve_placeholder(ph)

            # Check if this placeholder also appears inside a loop
            lp = ''
            for loop_name, children in loop_children.items():
                if ph in children:
                    lp = loop_name
                    break

            if lp:
                # Placeholder exists in BOTH top-level and loop context.
                # Create a top-level mapping first.
                self.env['sm.report.template.field'].create({
                    'template_id': self.id,
                    'placeholder': ph,
                    'field_path': field_path,
                    'loop_parent': False,
                    'auto_resolved': resolved,
                    'sequence': seq,
                })
                # Then create a loop-child mapping
                seq += 10
                child_path, child_resolved = (
                    self._auto_resolve_loop_child(lp, ph)
                    if not resolved else (field_path, resolved)
                )
                self.env['sm.report.template.field'].create({
                    'template_id': self.id,
                    'placeholder': ph,
                    'field_path': child_path,
                    'loop_parent': lp,
                    'auto_resolved': child_resolved,
                    'sequence': seq,
                })
            else:
                self.env['sm.report.template.field'].create({
                    'template_id': self.id,
                    'placeholder': ph,
                    'field_path': field_path,
                    'loop_parent': False,
                    'auto_resolved': resolved,
                    'sequence': seq,
                })

        # Create mappings for placeholders that ONLY appear inside loops
        for loop_name, children in loop_children.items():
            for ph in sorted(children - simple):
                seq += 10
                child_path, child_resolved = self._auto_resolve_loop_child(
                    loop_name, ph,
                )
                self.env['sm.report.template.field'].create({
                    'template_id': self.id,
                    'placeholder': ph,
                    'field_path': child_path,
                    'loop_parent': loop_name,
                    'auto_resolved': child_resolved,
                    'sequence': seq,
                })

        self.state = 'ready'
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Template Parsed'),
                'message': _(
                    'Found %(ph)s placeholders and %(lp)s loops.',
                    ph=len(simple), lp=len(loops),
                ),
                'type': 'success',
                'sticky': False,
            },
        }

    def action_activate(self):
        """Create a server action so users can print from the model."""
        self.ensure_one()
        unmapped = self.field_ids.filtered(
            lambda f: not f.field_path and not f.is_loop
        )
        if unmapped:
            raise UserError(_(
                'These placeholders are not mapped yet: %(names)s',
                names=', '.join(unmapped.mapped('placeholder')),
            ))
        unmapped_loops = self.field_ids.filtered(
            lambda f: f.is_loop and not f.field_path
        )
        if unmapped_loops:
            raise UserError(_(
                'These loop markers need a field path (One2many): %(names)s',
                names=', '.join(unmapped_loops.mapped('placeholder')),
            ))

        if not self.server_action_id:
            action = self.env['ir.actions.server'].sudo().create({
                'name': _('Print: %s') % self.name,
                'model_id': self.model_id.id,
                'binding_model_id': self.model_id.id,
                'binding_type': 'action',
                'state': 'code',
                'code': (
                    'action = env["sm.report.template"]'
                    '.browse(%d)'
                    '.generate_report(records)' % self.id
                ),
            })
            self.server_action_id = action

        self.state = 'active'

    def action_deactivate(self):
        """Remove the print action."""
        self.ensure_one()
        if self.server_action_id:
            self.server_action_id.sudo().unlink()
        self.state = 'ready'

    def action_reset_draft(self):
        """Reset to draft, removing action and mappings."""
        self.ensure_one()
        if self.server_action_id:
            self.server_action_id.sudo().unlink()
        self.field_ids.unlink()
        self.state = 'draft'

    def action_test_report(self):
        """Generate a test report with the first record of the model."""
        self.ensure_one()
        if not self.model_name:
            raise UserError(_('Please select a model first.'))
        Model = self.env[self.model_name]
        record = Model.search([], limit=1)
        if not record:
            raise UserError(_(
                'No records found in %(model)s to test with.',
                model=self.model_name,
            ))
        return self.generate_report(record)

    # ------------------------------------------------------------------
    #  Report generation (called from server action)
    # ------------------------------------------------------------------

    def generate_report(self, records):
        """Generate the report for *records* and return a download action."""
        self.ensure_one()
        if not records:
            raise UserError(_('No records selected.'))

        # Accept both recordsets and plain ID lists (e.g. from XML-RPC / server action)
        if isinstance(records, (list, tuple)):
            records = self.env[self.model_name].browse(records)
        elif isinstance(records, int):
            records = self.env[self.model_name].browse([records])

        file_bytes = base64.b64decode(self.template_file)
        outputs = []

        for record in records:
            ctx = self._build_render_context(record)

            if self.template_type == 'docx':
                output = self._render_docx(file_bytes, ctx)
            elif self.template_type == 'xlsx':
                output = self._render_xlsx(file_bytes, ctx)
            else:
                raise UserError(_('Unsupported template type.'))

            if self.output_format not in ('original', False):
                output = self._convert_format(
                    output, self.template_type, self.output_format,
                )

            outputs.append((record, output))

        if self.output_format in ('original', False):
            ext = self.template_type
        else:
            ext = self.output_format

        # --- Cover page (PDF only) ---
        cover_pdf = None
        if self.cover_file and ext == 'pdf':
            cover_bytes = base64.b64decode(self.cover_file)
            cover_ctx = self._build_render_context(records[0])
            cover_output = self._render_docx(cover_bytes, cover_ctx)
            cover_pdf = self._convert_format(cover_output, 'docx', 'pdf')

        if len(outputs) == 1:
            record, data = outputs[0]
            if cover_pdf:
                data = self._merge_pdfs([cover_pdf, data])
            filename = '%s - %s.%s' % (
                self.name, record.display_name or record.id, ext,
            )
            attachment = self.env['ir.attachment'].create({
                'name': filename,
                'datas': base64.b64encode(data),
                'res_model': self.model_name,
                'res_id': record.id,
                'type': 'binary',
            })
            return {
                'type': 'ir.actions.act_url',
                'url': '/web/content/%d?download=true' % attachment.id,
                'target': 'self',
            }

        # Multiple records — merge PDF or fall back to ZIP
        if ext == 'pdf' and self.merge_pdf:
            pdfs = [data for _rec, data in outputs]
            if cover_pdf:
                pdfs.insert(0, cover_pdf)
            merged = self._merge_pdfs(pdfs)
            attachment = self.env['ir.attachment'].create({
                'name': '%s.pdf' % self.name,
                'datas': base64.b64encode(merged),
                'type': 'binary',
            })
            return {
                'type': 'ir.actions.act_url',
                'url': '/web/content/%d?download=true' % attachment.id,
                'target': 'self',
            }

        # Fall back: ZIP
        zip_buf = BytesIO()
        with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
            for record, data in outputs:
                fname = '%s - %s.%s' % (
                    self.name, record.display_name or record.id, ext,
                )
                zf.writestr(fname, data)
        attachment = self.env['ir.attachment'].create({
            'name': '%s.zip' % self.name,
            'datas': base64.b64encode(zip_buf.getvalue()),
            'type': 'binary',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'self',
        }

    # ------------------------------------------------------------------
    #  Context building
    # ------------------------------------------------------------------

    def _build_render_context(self, record):
        """Build a flat {placeholder: value} dict plus loop data."""
        context = {}
        raw = {}
        loops = {}

        for fm in self.field_ids:
            if fm.is_loop:
                items = self._resolve_field_path(record, fm.field_path)
                children = self.field_ids.filtered(
                    lambda f, lp=fm.placeholder: (
                        f.loop_parent == lp and not f.is_loop
                    )
                )
                loop_items = []
                if hasattr(items, '__iter__') and not isinstance(items, str):
                    for item in items:
                        item_ctx = {}
                        item_raw = {}
                        for child in children:
                            val = self._resolve_field_path(item, child.field_path)
                            item_raw[child.placeholder] = val
                            item_ctx[child.placeholder] = self._format_value(val)
                        item_ctx['__raw__'] = item_raw
                        item_ctx['__record__'] = item
                        loop_items.append(item_ctx)
                loops[fm.placeholder] = loop_items

            elif not fm.loop_parent:
                val = self._resolve_field_path(record, fm.field_path)
                raw[fm.placeholder] = val
                context[fm.placeholder] = self._format_value(val)

        context['__loops__'] = loops
        context['__raw__'] = raw
        context['__record__'] = record
        return context

    # ------------------------------------------------------------------
    #  Field resolution helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _resolve_field_path(record, field_path):
        """Safely walk a dot-separated field path on a record."""
        if not field_path:
            return ''
        obj = record
        for part in field_path.split('.'):
            if isinstance(obj, models.BaseModel):
                if part not in obj._fields:
                    _logger.warning(
                        'Field %s not found on model %s', part, obj._name,
                    )
                    return ''
                try:
                    obj = obj[part]
                except Exception:
                    return ''
            else:
                return str(obj) if obj else ''
        return obj

    @staticmethod
    def _format_value(value):
        """Convert a field value to a display string."""
        if isinstance(value, bool):
            return _('Yes') if value else _('No')
        if isinstance(value, models.BaseModel):
            return value.display_name or ''
        if value is False or value is None:
            return ''
        return str(value)

    # ------------------------------------------------------------------
    #  Placeholder resolution (expressions, pipes, images)
    # ------------------------------------------------------------------

    def _resolve_placeholder(self, ph, context):
        """Resolve a single placeholder to its display value.

        Handles:
        - ``{{field}}``          → simple lookup
        - ``{{=expr}}``          → Python expression via safe_eval
        - ``{{field|format}}``   → value + format pipe
        - ``{{#...}}`` / ``{{/...}}`` → control tags → ''
        - ``{{img:field}}``      → '' (images handled in dedicated pass)
        """
        if ph.startswith('#') or ph.startswith('/'):
            return ''
        if ph.startswith('='):
            return str(self._eval_expression(ph[1:], context))
        if ph.startswith('img:'):
            return ''
        if '|' in ph:
            field_name, fmt = ph.split('|', 1)
            field_name = field_name.strip()
            fmt = fmt.strip()
            record = context.get('__record__')
            raw_val = context.get('__raw__', {}).get(field_name)
            if raw_val is None and record:
                raw_val = self._resolve_field_path(record, field_name)
            return self._apply_format_pipe(raw_val, fmt, record)
        return str(context.get(ph, '{{%s}}' % ph))

    def _eval_expression(self, expr, context):
        """Evaluate ``{{=expr}}`` using safe_eval."""
        import time as _time  # noqa: PLC0415
        import datetime as _datetime  # noqa: PLC0415

        record = context.get('__record__')
        eval_ctx = {
            'o': record,
            'record': record,
            'env': self.env,
            'user': self.env.user,
            'company_id': self.env.company,
            'time': _time,
            'datetime': _datetime,
        }
        # Also expose raw field values as local variables
        raw = context.get('__raw__', {})
        eval_ctx.update(raw)
        try:
            result = safe_eval(expr.strip(), eval_ctx)
            if result is None:
                return ''
            return self._format_value(result)
        except Exception as e:
            _logger.warning('Expression error «%s»: %s', expr, e)
            return '{{=%s}}' % expr

    def _eval_condition(self, condition, context):
        """Evaluate a condition string for ``{{#if …}}``."""
        import time as _time  # noqa: PLC0415
        import datetime as _datetime  # noqa: PLC0415

        record = context.get('__record__')

        # Simple field truthiness: {{#if field_name}}
        if re.match(r'^[a-zA-Z_][a-zA-Z0-9_.]*$', condition):
            val = context.get(condition)
            if val is None:
                raw = context.get('__raw__', {})
                val = raw.get(condition)
            if val is None and record:
                val = self._resolve_field_path(record, condition)
            return bool(val) and val != '' and val != _('No')

        # Full expression: {{#if amount_total > 1000}}
        eval_ctx = {
            'o': record,
            'record': record,
            'env': self.env,
            'user': self.env.user,
            'time': _time,
            'datetime': _datetime,
        }
        raw = context.get('__raw__', {})
        eval_ctx.update(raw)
        try:
            return bool(safe_eval(condition, eval_ctx))
        except Exception:
            _logger.warning('Condition eval failed: %s', condition)
            return False

    def _apply_format_pipe(self, raw_value, fmt, record=None):
        """Apply a format pipe to a raw field value.

        Supported pipes: ``date``, ``datetime``, ``currency``, ``number``,
        ``upper``, ``lower``, ``title``.
        """
        import datetime as _dt  # noqa: PLC0415

        if raw_value is False or raw_value is None:
            return ''
        fmt = fmt.lower()

        if fmt == 'upper':
            return str(raw_value).upper()
        if fmt == 'lower':
            return str(raw_value).lower()
        if fmt == 'title':
            return str(raw_value).title()

        if fmt == 'date':
            if isinstance(raw_value, (_dt.date, _dt.datetime)):
                if record and format_date:
                    return format_date(record.env, raw_value)
                return raw_value.strftime('%d/%m/%Y')
            return str(raw_value)

        if fmt == 'datetime':
            if isinstance(raw_value, _dt.datetime):
                if record and format_datetime:
                    return format_datetime(record.env, raw_value)
                return raw_value.strftime('%d/%m/%Y %H:%M')
            return str(raw_value)

        if fmt == 'currency':
            if isinstance(raw_value, (int, float)):
                if record and formatLang:
                    currency = record.env.company.currency_id
                    return formatLang(record.env, raw_value,
                                     currency_obj=currency)
                return '{:,.2f}'.format(raw_value)
            return str(raw_value)

        if fmt == 'number':
            if isinstance(raw_value, (int, float)):
                if record and formatLang:
                    return formatLang(record.env, raw_value)
                return '{:,.2f}'.format(raw_value)
            return str(raw_value)

        return str(raw_value)

    # ------------------------------------------------------------------
    #  Image helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _detect_image_format(data):
        """Detect image format from binary header."""
        if data[:8] == b'\x89PNG\r\n\x1a\n':
            return 'png'
        if data[:2] == b'\xff\xd8':
            return 'jpeg'
        if data[:4] == b'GIF8':
            return 'gif'
        if data[:4] == b'RIFF' and len(data) > 11 and data[8:12] == b'WEBP':
            return 'webp'
        if data[:2] == b'BM':
            return 'bmp'
        return 'png'

    @staticmethod
    def _get_image_size_emu(image_bytes,
                            max_width_cm=15, max_height_cm=10):
        """Return (cx, cy) in EMU, scaled to fit max dimensions."""
        EMU_PER_CM = 360000
        default_cx = 5 * EMU_PER_CM
        default_cy = 5 * EMU_PER_CM

        w_px = h_px = None

        if PILImage is not None:
            try:
                img = PILImage.open(BytesIO(image_bytes))
                w_px, h_px = img.size
            except Exception:
                pass
        else:
            # Fallback: read PNG / JPEG header with struct
            try:
                if image_bytes[:8] == b'\x89PNG\r\n\x1a\n':
                    w_px = struct.unpack('>I', image_bytes[16:20])[0]
                    h_px = struct.unpack('>I', image_bytes[20:24])[0]
                elif image_bytes[:2] == b'\xff\xd8':
                    # JPEG — scan for SOF marker
                    offset = 2
                    while offset < len(image_bytes) - 9:
                        marker = struct.unpack('>H', image_bytes[offset:offset+2])[0]
                        if 0xFFC0 <= marker <= 0xFFC3:
                            h_px = struct.unpack(
                                '>H', image_bytes[offset+5:offset+7])[0]
                            w_px = struct.unpack(
                                '>H', image_bytes[offset+7:offset+9])[0]
                            break
                        length = struct.unpack(
                            '>H', image_bytes[offset+2:offset+4])[0]
                        offset += 2 + length
            except Exception:
                pass

        if not w_px or not h_px:
            return default_cx, default_cy

        dpi = 96
        w_cm = w_px / dpi * 2.54
        h_cm = h_px / dpi * 2.54

        if w_cm > max_width_cm:
            ratio = max_width_cm / w_cm
            w_cm *= ratio
            h_cm *= ratio
        if h_cm > max_height_cm:
            ratio = max_height_cm / h_cm
            w_cm *= ratio
            h_cm *= ratio

        return int(w_cm * EMU_PER_CM), int(h_cm * EMU_PER_CM)

    # ------------------------------------------------------------------
    #  Placeholder parsing
    # ------------------------------------------------------------------

    def _parse_docx_placeholders(self, file_bytes):
        """Return (simple_set, loops_set, {loop: children_set}) from DOCX."""
        simple = set()
        loops = set()
        loop_children = {}
        current_loop = None

        with zipfile.ZipFile(BytesIO(file_bytes), 'r') as zf:
            for name in zf.namelist():
                if not (name.startswith('word/') and name.endswith('.xml')):
                    continue
                try:
                    tree = etree.fromstring(zf.read(name))
                except etree.XMLSyntaxError:
                    continue

                for para in tree.iter('{%s}p' % W_NS):
                    text = self._get_paragraph_text(para)
                    for match in re.finditer(r'\{\{(.+?)\}\}', text):
                        ph = match.group(1).strip()
                        if ph.startswith('=') or ph.startswith('img:'):
                            continue  # expressions & images resolved at render
                        if ph.startswith('#'):
                            tag = ph[1:]
                            if tag.startswith('if ') or tag == 'else':
                                continue  # conditionals
                            loop_name = tag
                            loops.add(loop_name)
                            loop_children.setdefault(loop_name, set())
                            current_loop = loop_name
                        elif ph.startswith('/'):
                            tag = ph[1:]
                            if tag == 'if':
                                continue
                            current_loop = None
                        else:
                            field_name = ph.split('|')[0].strip()
                            if current_loop:
                                loop_children[current_loop].add(field_name)
                            else:
                                simple.add(field_name)

        return simple, loops, loop_children

    def _parse_xlsx_placeholders(self, file_bytes):
        """Return (simple_set, loops_set, {loop: children_set}) from XLSX."""
        import openpyxl  # noqa: PLC0415  (bundled with Odoo)

        simple = set()
        loops = set()
        loop_children = {}
        current_loop = None

        wb = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if not (cell.value and isinstance(cell.value, str)):
                        continue
                    for match in re.finditer(r'\{\{(.+?)\}\}', cell.value):
                        ph = match.group(1).strip()
                        if ph.startswith('=') or ph.startswith('img:'):
                            continue
                        if ph.startswith('#'):
                            tag = ph[1:]
                            if tag.startswith('if ') or tag == 'else':
                                continue
                            loop_name = tag
                            loops.add(loop_name)
                            loop_children.setdefault(loop_name, set())
                            current_loop = loop_name
                        elif ph.startswith('/'):
                            tag = ph[1:]
                            if tag == 'if':
                                continue
                            current_loop = None
                        else:
                            field_name = ph.split('|')[0].strip()
                            if current_loop:
                                loop_children[current_loop].add(field_name)
                            else:
                                simple.add(field_name)
        wb.close()
        return simple, loops, loop_children

    # ------------------------------------------------------------------
    #  Auto-resolution helpers
    # ------------------------------------------------------------------

    def _auto_resolve_placeholder(self, placeholder):
        """Try to map *placeholder* to a real field path. Return (path, bool)."""
        if not self.model_name:
            return '', False
        Model = self.env[self.model_name]

        # Direct match
        if placeholder in Model._fields:
            return placeholder, True

        # Dot-notation check
        if '.' in placeholder:
            parts = placeholder.split('.')
            obj_model = Model
            valid = True
            for i, part in enumerate(parts):
                if part not in obj_model._fields:
                    valid = False
                    break
                field_def = obj_model._fields[part]
                if i < len(parts) - 1:
                    comodel = getattr(field_def, 'comodel_name', None)
                    if comodel:
                        obj_model = self.env[comodel]
                    else:
                        valid = False
                        break
            if valid:
                return placeholder, True

        return '', False

    def _auto_resolve_loop_child(self, loop_placeholder, child_placeholder):
        """Resolve *child_placeholder* against the comodel of the loop field."""
        if not self.model_name:
            return '', False

        # Find the loop's field path
        loop_map = self.field_ids.filtered(
            lambda f: f.placeholder == loop_placeholder and f.is_loop
        )
        loop_path = loop_map.field_path if loop_map else loop_placeholder

        # Walk the loop path to find the comodel
        Model = self.env[self.model_name]
        for part in (loop_path or loop_placeholder).split('.'):
            if part not in Model._fields:
                return '', False
            field_def = Model._fields[part]
            comodel = getattr(field_def, 'comodel_name', None)
            if comodel:
                Model = self.env[comodel]
            else:
                return '', False

        # Now resolve child in the comodel
        if child_placeholder in Model._fields:
            return child_placeholder, True

        if '.' in child_placeholder:
            parts = child_placeholder.split('.')
            obj = Model
            valid = True
            for i, p in enumerate(parts):
                if p not in obj._fields:
                    valid = False
                    break
                fd = obj._fields[p]
                if i < len(parts) - 1:
                    cm = getattr(fd, 'comodel_name', None)
                    if cm:
                        obj = self.env[cm]
                    else:
                        valid = False
                        break
            if valid:
                return child_placeholder, True

        return '', False

    # ==================================================================
    #  DOCX ENGINE
    # ==================================================================

    @staticmethod
    def _get_paragraph_text(para):
        """Concatenate all <w:t> texts in a paragraph."""
        parts = []
        for t in para.iter('{%s}t' % W_NS):
            if t.text:
                parts.append(t.text)
        return ''.join(parts)

    @staticmethod
    def _get_row_text(row):
        """Concatenate all paragraph texts in a table row."""
        parts = []
        for para in row.iter('{%s}p' % W_NS):
            parts.append(SmReportTemplate._get_paragraph_text(para))
        return ' '.join(parts)

    # ---- rendering ----

    def _render_docx(self, file_bytes, context):
        """Render a DOCX template, return bytes of the filled document."""
        buf_in = BytesIO(file_bytes)
        images = {}       # {rId: (filename, image_bytes)}
        img_counter = [1000]  # mutable counter for unique IDs

        # First pass: process XML parts (may populate images dict)
        processed = {}
        with zipfile.ZipFile(buf_in, 'r') as zin:
            infos = zin.infolist()
            for item in infos:
                data = zin.read(item.filename)
                if (
                    item.filename.startswith('word/')
                    and item.filename.endswith('.xml')
                ):
                    try:
                        data = self._process_docx_xml(
                            data, context, images, img_counter,
                        )
                    except Exception:
                        _logger.exception(
                            'Error processing %s', item.filename,
                        )
                processed[item.filename] = (item, data)

        # Second pass: update rels/content-types now that images are known
        buf_out = BytesIO()
        with zipfile.ZipFile(buf_out, 'w', zipfile.ZIP_DEFLATED) as zout:
            for filename, (item, data) in processed.items():
                if images:
                    if filename == 'word/_rels/document.xml.rels':
                        data = self._update_docx_rels(data, images)
                    elif filename == '[Content_Types].xml':
                        data = self._update_content_types(data, images)
                zout.writestr(item, data)

            # Write collected images into word/media/
            for rid, (fname, img_bytes) in images.items():
                zout.writestr('word/media/' + fname, img_bytes)

        return buf_out.getvalue()

    def _process_docx_xml(self, xml_bytes, context,
                          images=None, img_counter=None):
        """Replace placeholders in one XML part of the DOCX."""
        tree = etree.fromstring(xml_bytes)

        # 1) conditionals  — {{#if}} / {{#else}} / {{/if}}
        self._process_docx_conditionals(tree, context)

        # 2) images — {{img:field}}
        if images is not None:
            self._process_docx_images(tree, context, images, img_counter)

        # 3) table-row loops
        loops = context.get('__loops__', {})
        if loops:
            self._docx_table_loops(tree, loops)

        # 4) simple placeholder replacement
        for para in tree.iter('{%s}p' % W_NS):
            self._docx_replace_paragraph(para, context)

        return etree.tostring(
            tree, xml_declaration=True, encoding='UTF-8', standalone=True,
        )

    def _docx_replace_paragraph(self, para, context):
        """Replace {{…}} in a single paragraph."""
        text_elements = list(para.iter('{%s}t' % W_NS))
        if not text_elements:
            return

        full_text = ''.join(t.text or '' for t in text_elements)
        if '{{' not in full_text:
            return

        def _repl(m):
            ph = m.group(1).strip()
            return self._resolve_placeholder(ph, context)

        new_text = re.sub(r'\{\{(.+?)\}\}', _repl, full_text)
        if new_text == full_text:
            return

        # Put everything in the first <w:t>, blank the rest
        text_elements[0].text = new_text
        text_elements[0].set('{%s}space' % XML_NS, 'preserve')
        for t in text_elements[1:]:
            t.text = ''

    # ---- conditionals ----

    def _process_docx_conditionals(self, parent, context):
        """Process ``{{#if}}``, ``{{#else}}``, ``{{/if}}`` blocks.

        Works on body, table cells, or any container with paragraphs
        and table rows as children.  Supports nesting.
        """
        stack = []   # list of bools: True = currently showing
        to_remove = []

        for child in list(parent):
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag

            # Get text content
            text = ''
            if tag == 'p':
                text = self._get_paragraph_text(child)
            elif tag == 'tr':
                text = self._get_row_text(child)

            # {{#if condition}}
            if_match = re.search(r'\{\{#if\s+(.+?)\}\}', text) if text else None
            if if_match:
                parent_showing = stack[-1] if stack else True
                if parent_showing:
                    result = self._eval_condition(
                        if_match.group(1).strip(), context)
                    stack.append(result)
                else:
                    stack.append(False)
                to_remove.append(child)
                continue

            # {{#else}}
            if '{{#else}}' in text:
                if stack:
                    parent_showing = stack[-2] if len(stack) > 1 else True
                    if parent_showing:
                        stack[-1] = not stack[-1]
                to_remove.append(child)
                continue

            # {{/if}}
            if '{{/if}}' in text:
                if stack:
                    stack.pop()
                to_remove.append(child)
                continue

            # If currently inside a false block, remove
            if stack and not stack[-1]:
                to_remove.append(child)
                continue

            # Recurse into containers
            if tag in ('body', 'tbl', 'tc', 'txbxContent'):
                self._process_docx_conditionals(child, context)
            elif tag == 'tr':
                for tc in child.findall('{%s}tc' % W_NS):
                    self._process_docx_conditionals(tc, context)

        for elem in to_remove:
            try:
                parent.remove(elem)
            except ValueError:
                pass

    # ---- images ----

    def _process_docx_images(self, tree, context, images, img_counter):
        """Replace ``{{img:field_path}}`` with inline images."""
        record = context.get('__record__')
        if not record:
            return

        for para in list(tree.iter('{%s}p' % W_NS)):
            text = self._get_paragraph_text(para)
            if '{{img:' not in text:
                continue

            matches = list(re.finditer(r'\{\{img:(.+?)\}\}', text))
            if not matches:
                continue

            # Resolve images
            resolved = []
            for m in matches:
                field_path = m.group(1).strip()
                raw_val = self._resolve_field_path(record, field_path)
                img_data = None
                if raw_val:
                    if isinstance(raw_val, str):
                        try:
                            img_data = base64.b64decode(raw_val)
                        except Exception:
                            pass
                    elif isinstance(raw_val, bytes):
                        img_data = raw_val

                if img_data and len(img_data) >= 8:
                    fmt = self._detect_image_format(img_data)
                    cx, cy = self._get_image_size_emu(img_data)
                    rid = 'rId%d' % img_counter[0]
                    filename = 'image%d.%s' % (img_counter[0], fmt)
                    img_counter[0] += 1
                    images[rid] = (filename, img_data)
                    resolved.append((m.start(), m.end(), rid, filename,
                                     cx, cy))
                else:
                    resolved.append((m.start(), m.end(),
                                     None, None, 0, 0))

            # Rebuild paragraph runs
            for r in list(para.findall('{%s}r' % W_NS)):
                para.remove(r)

            last_end = 0
            for start, end, rid, filename, cx, cy in resolved:
                seg = text[last_end:start]
                if seg:
                    run = etree.SubElement(para, '{%s}r' % W_NS)
                    t_el = etree.SubElement(run, '{%s}t' % W_NS)
                    t_el.text = seg
                    t_el.set('{%s}space' % XML_NS, 'preserve')
                if rid:
                    img_run = self._create_image_run(
                        rid, filename, cx, cy, img_counter[0])
                    para.append(img_run)
                last_end = end

            tail = text[last_end:]
            if tail:
                run = etree.SubElement(para, '{%s}r' % W_NS)
                t_el = etree.SubElement(run, '{%s}t' % W_NS)
                t_el.text = tail
                t_el.set('{%s}space' % XML_NS, 'preserve')

    @staticmethod
    def _create_image_run(rid, filename, cx, cy, doc_id):
        """Build a ``<w:r>`` element containing an inline drawing."""
        xml = (
            '<w:r xmlns:w="{w}" xmlns:wp="{wp}" xmlns:a="{a}" '
            'xmlns:pic="{pic}" xmlns:r="{r}">'
            '<w:drawing>'
            '<wp:inline distT="0" distB="0" distL="0" distR="0">'
            '<wp:extent cx="{cx}" cy="{cy}"/>'
            '<wp:docPr id="{did}" name="Image {did}"/>'
            '<a:graphic>'
            '<a:graphicData uri="{pic}">'
            '<pic:pic>'
            '<pic:nvPicPr>'
            '<pic:cNvPr id="0" name="{fn}"/>'
            '<pic:cNvPicPr/>'
            '</pic:nvPicPr>'
            '<pic:blipFill>'
            '<a:blip r:embed="{rid}"/>'
            '<a:stretch><a:fillRect/></a:stretch>'
            '</pic:blipFill>'
            '<pic:spPr>'
            '<a:xfrm>'
            '<a:off x="0" y="0"/>'
            '<a:ext cx="{cx}" cy="{cy}"/>'
            '</a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            '</pic:spPr>'
            '</pic:pic>'
            '</a:graphicData>'
            '</a:graphic>'
            '</wp:inline>'
            '</w:drawing>'
            '</w:r>'
        ).format(
            w=W_NS, wp=WP_NS, a=A_NS, pic=PIC_NS, r=R_NS,
            cx=cx, cy=cy, did=doc_id, fn=filename, rid=rid,
        )
        return etree.fromstring(xml)

    @staticmethod
    def _update_docx_rels(rels_bytes, images):
        """Add image relationships to ``document.xml.rels``."""
        tree = etree.fromstring(rels_bytes)
        for rid, (filename, _data) in images.items():
            rel = etree.SubElement(tree, '{%s}Relationship' % REL_NS)
            rel.set('Id', rid)
            rel.set('Type', IMG_REL_TYPE)
            rel.set('Target', 'media/' + filename)
        return etree.tostring(
            tree, xml_declaration=True, encoding='UTF-8', standalone=True)

    @staticmethod
    def _update_content_types(ct_bytes, images):
        """Ensure image content types are registered."""
        tree = etree.fromstring(ct_bytes)
        existing = {c.get('Extension', '').lower()
                     for c in tree if c.get('Extension')}
        mime_map = {
            'png': 'image/png', 'jpeg': 'image/jpeg', 'jpg': 'image/jpeg',
            'gif': 'image/gif', 'bmp': 'image/bmp', 'webp': 'image/webp',
        }
        needed = set()
        for _rid, (fn, _data) in images.items():
            ext = fn.rsplit('.', 1)[-1].lower()
            needed.add(ext)
        for ext in needed - existing:
            ct = mime_map.get(ext, 'application/octet-stream')
            el = etree.SubElement(tree, '{%s}Default' % CT_NS)
            el.set('Extension', ext)
            el.set('ContentType', ct)
        return etree.tostring(
            tree, xml_declaration=True, encoding='UTF-8', standalone=True)

    # ---- table-row loops ----

    def _docx_table_loops(self, tree, loops):
        """Expand {{#name}} / {{/name}} table row blocks."""
        for table in tree.iter('{%s}tbl' % W_NS):
            self._docx_expand_table(table, loops)

    def _docx_expand_table(self, table, loops):
        """Process one <w:tbl> for loop expansion. Restarts on mutation."""
        changed = True
        while changed:
            changed = False
            rows = list(table.findall('{%s}tr' % W_NS))

            for loop_name, items in loops.items():
                start_tag = '{{#%s}}' % loop_name
                end_tag = '{{/%s}}' % loop_name
                start_idx = end_idx = None

                for i, row in enumerate(rows):
                    txt = self._get_row_text(row)
                    if start_tag in txt and start_idx is None:
                        start_idx = i
                    if end_tag in txt:
                        end_idx = i
                        break

                if start_idx is None or end_idx is None or end_idx <= start_idx:
                    continue

                # Rows between markers are the template
                template_rows = rows[start_idx + 1:end_idx]

                # Remember the element after the end-marker for insertion
                after = rows[end_idx + 1] if end_idx + 1 < len(rows) else None

                # Remove marker rows and template rows
                table.remove(rows[start_idx])
                for tr in template_rows:
                    table.remove(tr)
                table.remove(rows[end_idx])

                # Insert expanded rows
                for item_ctx in items:
                    for tmpl_row in template_rows:
                        new_row = deepcopy(tmpl_row)
                        for para in new_row.iter('{%s}p' % W_NS):
                            self._docx_replace_paragraph(para, item_ctx)
                        if after is not None:
                            after.addprevious(new_row)
                        else:
                            table.append(new_row)

                changed = True
                break  # restart — row indices are stale

    # ==================================================================
    #  XLSX ENGINE
    # ==================================================================

    def _render_xlsx(self, file_bytes, context):
        """Render an XLSX template, return bytes."""
        import openpyxl  # noqa: PLC0415
        from copy import copy  # noqa: PLC0415

        wb = openpyxl.load_workbook(BytesIO(file_bytes))

        loops = context.get('__loops__', {})

        for ws in wb.worksheets:
            # 1) conditionals
            self._process_xlsx_conditionals(ws, context)
            # 2) expand loops
            if loops:
                self._xlsx_expand_loops(ws, loops, copy)
            # 3) images
            self._process_xlsx_images(ws, context)
            # 4) simple replacement
            self._xlsx_replace_cells(ws, context)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _xlsx_replace_cells(self, ws, context):
        """Replace {{…}} in every cell of a worksheet."""
        for row in ws.iter_rows():
            for cell in row:
                if not (cell.value and isinstance(cell.value, str)):
                    continue
                if '{{' not in cell.value:
                    continue

                def _repl(m, _ctx=context, _self=self):
                    ph = m.group(1).strip()
                    return _self._resolve_placeholder(ph, _ctx)

                cell.value = re.sub(r'\{\{(.+?)\}\}', _repl, cell.value)

    def _xlsx_expand_loops(self, ws, loops, copy_fn):
        """Expand loop rows in a worksheet."""
        for loop_name, items in loops.items():
            if not items:
                continue
            start_tag = '{{#%s}}' % loop_name
            end_tag = '{{/%s}}' % loop_name
            start_row = end_row = None

            for row in ws.iter_rows(min_row=1):
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        if start_tag in cell.value:
                            start_row = cell.row
                        if end_tag in cell.value:
                            end_row = cell.row

            if not start_row or not end_row or end_row <= start_row:
                continue

            tmpl_start = start_row + 1
            tmpl_end = end_row - 1
            num_tmpl = tmpl_end - tmpl_start + 1
            if num_tmpl <= 0:
                continue

            # Store template row data (values + styles)
            template_data = []
            for r in range(tmpl_start, tmpl_end + 1):
                row_data = []
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(row=r, column=c)
                    row_data.append({
                        'value': cell.value,
                        'font': copy_fn(cell.font),
                        'border': copy_fn(cell.border),
                        'fill': copy_fn(cell.fill),
                        'number_format': cell.number_format,
                        'alignment': copy_fn(cell.alignment),
                    })
                template_data.append(row_data)

            # Clear marker cells
            for cell in ws[start_row]:
                if cell.value and isinstance(cell.value, str):
                    cell.value = cell.value.replace(start_tag, '').strip() or None
            for cell in ws[end_row]:
                if cell.value and isinstance(cell.value, str):
                    cell.value = cell.value.replace(end_tag, '').strip() or None

            # Delete original template rows
            ws.delete_rows(tmpl_start, num_tmpl)
            # Adjust end_row position
            end_row -= num_tmpl
            # Delete marker rows (end first, then start)
            ws.delete_rows(end_row)
            ws.delete_rows(start_row)

            # Insert new rows for all items
            total_new = len(items) * num_tmpl
            ws.insert_rows(start_row, total_new)

            # Fill data
            cur_row = start_row
            for item_ctx in items:
                for tmpl_row_data in template_data:
                    for c_idx, cd in enumerate(tmpl_row_data, start=1):
                        cell = ws.cell(row=cur_row, column=c_idx)
                        value = cd['value']
                        if value and isinstance(value, str) and '{{' in value:
                            value = re.sub(
                                r'\{\{(.+?)\}\}',
                                lambda m, _c=item_ctx: str(
                                    _c.get(m.group(1).strip(), m.group(0))
                                ),
                                value,
                            )
                        cell.value = value
                        cell.font = cd['font']
                        cell.border = cd['border']
                        cell.fill = cd['fill']
                        cell.number_format = cd['number_format']
                        cell.alignment = cd['alignment']
                    cur_row += 1

    # ---- XLSX conditionals ----

    def _process_xlsx_conditionals(self, ws, context):
        """Remove rows inside false ``{{#if}}`` blocks."""
        rows_to_delete = []
        stack = []  # list of bools

        for row in ws.iter_rows(min_row=1):
            row_num = row[0].row
            row_text = ' '.join(
                str(c.value) for c in row if c.value
            )

            if_match = re.search(r'\{\{#if\s+(.+?)\}\}', row_text)
            if if_match:
                parent_showing = stack[-1] if stack else True
                if parent_showing:
                    result = self._eval_condition(
                        if_match.group(1).strip(), context)
                    stack.append(result)
                else:
                    stack.append(False)
                rows_to_delete.append(row_num)
                continue

            if '{{#else}}' in row_text:
                if stack:
                    parent_showing = stack[-2] if len(stack) > 1 else True
                    if parent_showing:
                        stack[-1] = not stack[-1]
                rows_to_delete.append(row_num)
                continue

            if '{{/if}}' in row_text:
                if stack:
                    stack.pop()
                rows_to_delete.append(row_num)
                continue

            if stack and not stack[-1]:
                rows_to_delete.append(row_num)

        # Delete from bottom up so indices stay valid
        for row_num in reversed(sorted(rows_to_delete)):
            ws.delete_rows(row_num)

    # ---- XLSX images ----

    def _process_xlsx_images(self, ws, context):
        """Replace ``{{img:field}}`` cells with images in XLSX."""
        record = context.get('__record__')
        if not record:
            return

        try:
            from openpyxl.drawing.image import Image as XlImage  # noqa: PLC0415
        except ImportError:
            return

        for row in list(ws.iter_rows()):
            for cell in row:
                if not (cell.value and isinstance(cell.value, str)):
                    continue
                m = re.search(r'\{\{img:(.+?)\}\}', cell.value)
                if not m:
                    continue

                field_path = m.group(1).strip()
                raw_val = self._resolve_field_path(record, field_path)
                img_data = None
                if raw_val:
                    if isinstance(raw_val, str):
                        try:
                            img_data = base64.b64decode(raw_val)
                        except Exception:
                            pass
                    elif isinstance(raw_val, bytes):
                        img_data = raw_val

                cell.value = ''
                if img_data and len(img_data) >= 8:
                    try:
                        xl_img = XlImage(BytesIO(img_data))
                        xl_img.anchor = cell.coordinate
                        ws.add_image(xl_img)
                    except Exception:
                        _logger.warning(
                            'Failed to insert XLSX image for %s', field_path)

    # ==================================================================
    #  PDF MERGE
    # ==================================================================

    @staticmethod
    def _merge_pdfs(pdf_list):
        """Merge a list of PDF byte-strings into a single PDF."""
        # Try modern pypdf first, then PyPDF2 (new API), then PyPDF2 (old API)
        PdfReader = PdfWriter = None
        legacy = False

        try:
            from pypdf import PdfReader, PdfWriter  # noqa: PLC0415
        except ImportError:
            try:
                from PyPDF2 import PdfReader, PdfWriter  # noqa: PLC0415
            except ImportError:
                try:
                    from PyPDF2 import PdfFileReader as PdfReader, PdfFileWriter as PdfWriter  # noqa: PLC0415,E501
                    legacy = True
                except ImportError:
                    raise UserError(_(
                        'PDF merge requires the "pypdf" library.\n\n'
                        'Install it with:  pip install pypdf'
                    ))

        writer = PdfWriter()
        for pdf_bytes in pdf_list:
            reader = PdfReader(BytesIO(pdf_bytes))
            if legacy:
                for i in range(reader.getNumPages()):
                    writer.addPage(reader.getPage(i))
            else:
                for page in reader.pages:
                    writer.add_page(page)

        buf = BytesIO()
        writer.write(buf)
        return buf.getvalue()

    # ==================================================================
    #  DOMAIN-BASED TEMPLATE SELECTION
    # ==================================================================

    @api.model
    def get_templates_for_record(self, record):
        """Return templates matching *record* based on model + domain rules.

        Templates without domain are always returned.
        Templates with domain are evaluated against the record.
        Results are sorted by priority (lowest first).
        """
        model_name = record._name
        templates = self.search([
            ('model_name', '=', model_name),
            ('state', '=', 'active'),
        ], order='priority, name')

        matched = self.browse()
        for tmpl in templates:
            if not tmpl.use_domain or not tmpl.template_domain \
                    or tmpl.template_domain in ('[]', ''):
                matched |= tmpl
            else:
                try:
                    domain = safe_eval(tmpl.template_domain)
                    if record.filtered_domain(domain):
                        matched |= tmpl
                except Exception:
                    _logger.warning(
                        'Invalid domain on template %s: %s',
                        tmpl.name, tmpl.template_domain,
                    )
        return matched

    @api.model
    def action_print_with_selection(self, res_model, res_ids):
        """Show a template picker when multiple templates match, or print
        directly when only one matches. Called from JS or server action."""
        records = self.env[res_model].browse(res_ids)
        if not records:
            raise UserError(_('No records selected.'))

        templates = self.get_templates_for_record(records[0])
        if not templates:
            raise UserError(_(
                'No active report templates found for %(model)s.',
                model=res_model,
            ))

        if len(templates) == 1:
            return templates.generate_report(records)

        # Multiple templates — show selection wizard
        return {
            'type': 'ir.actions.act_window',
            'name': _('Select Report Template'),
            'res_model': 'sm.report.template.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {
                'default_res_model': res_model,
                'default_res_ids': res_ids,
                'default_available_template_ids': templates.ids,
            },
        }

    # ==================================================================
    #  PDF PREVIEW
    # ==================================================================

    def action_preview_report(self):
        """Generate a preview PDF and open it in a new browser tab."""
        self.ensure_one()
        if not self.model_name:
            raise UserError(_('Please select a model first.'))
        Model = self.env[self.model_name]
        record = Model.search([], limit=1)
        if not record:
            raise UserError(_(
                'No records found in %(model)s to preview.',
                model=self.model_name,
            ))

        file_bytes = base64.b64decode(self.template_file)
        ctx = self._build_render_context(record)

        if self.template_type == 'docx':
            output = self._render_docx(file_bytes, ctx)
        elif self.template_type == 'xlsx':
            output = self._render_xlsx(file_bytes, ctx)
        else:
            raise UserError(_('Unsupported template type.'))

        # Always convert to PDF for preview
        pdf_bytes = self._convert_format(output, self.template_type, 'pdf')

        attachment = self.env['ir.attachment'].create({
            'name': 'Preview - %s.pdf' % self.name,
            'datas': base64.b64encode(pdf_bytes),
            'res_model': self._name,
            'res_id': self.id,
            'type': 'binary',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=false' % attachment.id,
            'target': 'new',
        }

    # ==================================================================
    #  PDF CONVERSION (cross-platform: Linux / macOS / Windows)
    # ==================================================================

    @staticmethod
    def _find_libreoffice():
        """Locate the LibreOffice binary on Linux, macOS, or Windows."""
        system = platform.system()

        if system == 'Linux':
            candidates = [
                'libreoffice',
                '/usr/bin/libreoffice',
                '/usr/lib/libreoffice/program/soffice',
                '/snap/bin/libreoffice',
                'soffice',
            ]
        elif system == 'Darwin':  # macOS
            candidates = [
                '/Applications/LibreOffice.app/Contents/MacOS/soffice',
                'libreoffice',
                'soffice',
            ]
        elif system == 'Windows':
            candidates = [
                r'C:\Program Files\LibreOffice\program\soffice.exe',
                r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
                'soffice.exe',
                'soffice',
            ]
        else:
            candidates = ['libreoffice', 'soffice']

        for cmd in candidates:
            found = shutil.which(cmd)
            if found:
                return found
            if os.path.isfile(cmd) and os.access(cmd, os.X_OK):
                return cmd

        return None

    @staticmethod
    def _get_install_hint():
        """Return OS-specific install instructions for LibreOffice."""
        system = platform.system()
        if system == 'Linux':
            return (
                'Linux:  sudo apt-get install libreoffice-core\n'
                '   or:  sudo dnf install libreoffice-core\n'
                '   or:  sudo snap install libreoffice'
            )
        if system == 'Darwin':
            return (
                'macOS:  brew install --cask libreoffice\n'
                '   or:  Download from https://www.libreoffice.org/download/'
            )
        if system == 'Windows':
            return (
                'Windows:  Download and install from '
                'https://www.libreoffice.org/download/\n'
                'Make sure soffice.exe is in your PATH or installed in the '
                'default location.'
            )
        return 'Install LibreOffice from https://www.libreoffice.org/download/'

    @staticmethod
    def _convert_to_pdf(file_bytes, source_type):
        """Convert DOCX/XLSX bytes to PDF. Backward-compatible wrapper."""
        return SmReportTemplate._convert_format(file_bytes, source_type, 'pdf')

    @staticmethod
    def _convert_format(file_bytes, source_type, target_format):
        """Convert DOCX/XLSX bytes to target format via LibreOffice."""
        lo_bin = SmReportTemplate._find_libreoffice()
        if not lo_bin:
            raise UserError(_(
                'LibreOffice is not installed on this server.\n\n%s'
            ) % SmReportTemplate._get_install_hint())

        tmp_in = tempfile.NamedTemporaryFile(
            suffix='.%s' % source_type, delete=False,
        )
        try:
            tmp_in.write(file_bytes)
            tmp_in.flush()
            tmp_in.close()

            out_dir = tempfile.mkdtemp()
            try:
                cmd = [
                    lo_bin, '--headless', '--norestore',
                    '--convert-to', target_format,
                    '--outdir', out_dir, tmp_in.name,
                ]

                # On Windows, hide the console window
                kwargs = {}
                if platform.system() == 'Windows':
                    si = subprocess.STARTUPINFO()
                    si.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                    kwargs['startupinfo'] = si

                result = subprocess.run(
                    cmd, capture_output=True, timeout=120, **kwargs,
                )
                if result.returncode != 0:
                    _logger.error(
                        'LibreOffice conversion failed:\n'
                        'Binary: %s\nOS: %s\nFormat: %s\n'
                        'stdout: %s\nstderr: %s',
                        lo_bin, platform.system(), target_format,
                        result.stdout.decode(errors='replace'),
                        result.stderr.decode(errors='replace'),
                    )
                    raise UserError(_(
                        '%(fmt)s conversion failed. LibreOffice returned an error.\n'
                        'Binary used: %(binary)s\n\n%(hint)s',
                        fmt=target_format.upper(),
                        binary=lo_bin,
                        hint=SmReportTemplate._get_install_hint(),
                    ))

                out_name = (
                    os.path.splitext(os.path.basename(tmp_in.name))[0]
                    + '.%s' % target_format
                )
                out_path = os.path.join(out_dir, out_name)
                if not os.path.isfile(out_path):
                    raise UserError(_(
                        '%(fmt)s conversion completed but output file was not found.\n'
                        'Expected: %(path)s\n\n'
                        'This may happen if LibreOffice is not fully installed.',
                        fmt=target_format.upper(),
                        path=out_path,
                    ))
                with open(out_path, 'rb') as f:
                    return f.read()
            finally:
                shutil.rmtree(out_dir, ignore_errors=True)
        except FileNotFoundError:
            raise UserError(_(
                'LibreOffice binary not found: %(binary)s\n\n%(hint)s',
                binary=lo_bin,
                hint=SmReportTemplate._get_install_hint(),
            ))
        finally:
            try:
                os.unlink(tmp_in.name)
            except OSError:
                pass

    # ------------------------------------------------------------------
    #  ORM overrides
    # ------------------------------------------------------------------

    def write(self, vals):
        if 'template_file' in vals:
            for rec in self:
                if rec.state == 'active':
                    rec.action_deactivate()
                rec.field_ids.unlink()
            vals['state'] = 'draft'
        return super().write(vals)

    def unlink(self):
        for rec in self:
            if rec.server_action_id:
                rec.server_action_id.sudo().unlink()
        return super(SmReportTemplate, self).unlink()
